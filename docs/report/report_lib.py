"""Builds the .docx report from the group's UIT report template (template/uit_report_template.docx).

Building blocks: chapter/section headings, paragraphs with inline formatting (**bold**, *italic*, `code`),
bullet lists, data tables (navy header), SQL code boxes with syntax highlighting, figures with captions,
table of contents / list of figures / list of tables (Word fields).
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1F3864"
ACCENT = "2E75B6"
BORDER = "BFBFBF"
CODE_TITLE_BG = "E7ECF5"
CODE_BG = "F7F7F7"
INLINE_CODE = "C7254E"
PAGE_WIDTH_TWIPS = 9070  # A4, 2.5 cm left/right margins
# A longer caption wraps to a second line in the list of figures: the longest one that still fit had 70
# characters, 60 leaves room for wide letters and two-digit figure numbers
FIGURE_CAPTION_MAX = 60

SQL_KEYWORDS = set("""
ADD AFTER ALL ALTER AND APPLY AS ASC AUTHORIZATION BACKUP BEGIN BETWEEN BREAK BY CASCADE CASE CATCH CHECK CLOSE
COLLATE COLUMN COMMIT CONSTRAINT CONTAINMENT CONTINUE CREATE CROSS CURSOR DATABASE DEALLOCATE DECLARE DEFAULT
DELETE DENY DESC DIFFERENTIAL DISK DISTINCT DROP ELSE END EXEC EXECUTE EXISTS FAST_FORWARD FETCH FOR FOREIGN
FROM FULL FUNCTION GO GRANT GROUP HAVING IDENTITY IF IN INDEX INIT INNER INSERT INSTEAD INTO IS JOIN KEY LEFT
LIKE LOCAL LOG MEMBER MOVE NEXT NOCOUNT NORECOVERY NOT NULL OF OFF ON OPEN OR ORDER OUTER OUTPUT OVER OWNER
PARTITION PASSWORD PERSISTED PIVOT PRIMARY PROCEDURE RECOVERY REFERENCES REPLACE RESTORE RETURN RETURNS REVERT
REVOKE RIGHT ROLE ROLLBACK ROWS SCHEMA SCHEMABINDING SELECT SEQUENCE SET TABLE THEN THROW TO TOP TRAN TRANSACTION
TRIGGER TRY UNION UNIQUE UPDATE USE USER VALUES VIEW WHEN WHERE WHILE WITH XACT_ABORT RAISERROR PRINT
INT VARCHAR NVARCHAR CHAR DATE DATETIME TIME DECIMAL TINYINT SMALLINT BIGINT BIT XML SYSNAME UNBOUNDED PRECEDING
""".split())
SQL_FUNCS = set("""
COUNT SUM AVG MIN MAX ISNULL COALESCE CAST CONVERT GETDATE DATEADD DATEDIFF YEAR MONTH DAY ROUND LEFT RIGHT LEN
LTRIM RTRIM REPLACE STUFF NULLIF ROW_NUMBER DENSE_RANK QUOTENAME USER_NAME ORIGINAL_LOGIN OBJECT_ID IS_MEMBER
DATABASE_PRINCIPAL_ID CHECKSUM ABS FORMAT DATEFROMPARTS SERVERPROPERTY CHARINDEX SUBSTRING
GETUTCDATE SYSDATETIMEOFFSET SWITCHOFFSET TODATETIMEOFFSET
""".split())


# ---------------------------------------------------------------------------- XML helpers
def _set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_cell_borders(cell, color: str = BORDER, sz: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _set_cell_margins(cell, top=40, left=100, bottom=40, right=100) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _set_cell_width(cell, twips: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for cu in tcPr.findall(qn("w:tcW")):
        tcPr.remove(cu)
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(twips))
    w.set(qn("w:type"), "dxa")
    tcPr.append(w)


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Inserts a Word field (TOC, PAGE...); Word fills it in when the document is opened or on F9."""
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    run._r.append(fld)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run._r.append(it)
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "separate")
    run._r.append(fld)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run._r.append(fld)


def _check_figure_caption(caption: str) -> None:
    """Stops the build on a caption that would take two lines in the list of figures."""
    if len(caption) > FIGURE_CAPTION_MAX:
        raise ValueError(f"Figure caption longer than {FIGURE_CAPTION_MAX} characters ({len(caption)}), it wraps in "
                         f"the list of figures - shorten it: {caption}")


def _para_format(p, after=120, line=276, first_line=284, align="both", before=0, keep_next=False):
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(before))
    sp.set(qn("w:after"), str(after))
    sp.set(qn("w:line"), str(line))
    sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)
    if first_line:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLine"), str(first_line))
        pPr.append(ind)
    if align:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        pPr.append(jc)
    if keep_next:
        pPr.insert(0, OxmlElement("w:keepNext"))


# ---------------------------------------------------------------------------- SQL highlighting
_TOKEN_RE = re.compile(
    r"(--[^\n]*)|(/\*.*?\*/)|(N?'(?:[^']|'')*')|(@@?\w+)|(\b\d+(?:\.\d+)?\b)|([A-Za-z_][\w$#]*)|(\s+)|(.)",
    re.S,
)


def sql_tokens(code: str):
    """Returns a list of (text, color) pairs using the SSMS colors."""
    out = []
    for m in _TOKEN_RE.finditer(code):
        cmt1, cmt2, s, var, num, word, ws, other = m.groups()
        if cmt1 or cmt2:
            out.append((cmt1 or cmt2, "008000"))
        elif s:
            out.append((s, "A31515"))
        elif var:
            out.append((var, "000000"))
        elif num:
            out.append((num, "098658"))
        elif word:
            up = word.upper()
            if up in SQL_KEYWORDS:
                out.append((word, "0000FF"))
            elif up in SQL_FUNCS:
                out.append((word, "C2185B"))
            else:
                out.append((word, "000000"))
        elif ws:
            out.append((ws, None))
        else:
            out.append((other, "808080"))
    return out


# ---------------------------------------------------------------------------- code extracted from the SQL files
def sql_object(sql_dir: Path, filename: str, name: str) -> str:
    """Returns the CREATE ... dbo.<name> statement of a SQL file (up to the next GO line)."""
    text = (sql_dir / filename).read_text(encoding="utf-8")
    m = re.search(rf"^CREATE\s+(?:PROCEDURE|FUNCTION|TRIGGER|VIEW|TABLE)\s+dbo\.{re.escape(name)}\b.*?(?=^GO\s*$)",
                  text, re.S | re.M | re.I)
    if not m:
        raise KeyError(f"{name} not found in {filename}")
    return m.group(0).rstrip()


def sql_block(sql_dir: Path, filename: str, start_marker: str, end_marker: str | None = None) -> str:
    """Returns the code between two markers (including the line that holds start_marker)."""
    text = (sql_dir / filename).read_text(encoding="utf-8")
    i = text.index(start_marker)
    i = text.rfind("\n", 0, i) + 1
    j = text.index(end_marker, i + len(start_marker)) if end_marker else len(text)
    return text[i:j].rstrip()


# ---------------------------------------------------------------------------- report class
class Report:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        self.chapter = 0
        self.fig = 0
        self.tbl = 0
        self.code_no = 0
        self._ensure_styles()

    # ----- style
    def _ensure_styles(self):
        styles = self.doc.styles
        names = {s.name for s in styles}
        if "TableCaption" not in names:
            st = styles.add_style("TableCaption", 1)  # paragraph
            st.base_style = styles["FigureCaption"]
            st.font.italic = True
            st.font.size = Pt(11)
            st.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.space_before = Pt(6)
            st.paragraph_format.space_after = Pt(4)
            st.paragraph_format.keep_with_next = True

    @property
    def body(self):
        return self.doc.element.body

    def _move_before_sect(self, element):
        """python-docx appends to the end of the body; keep sectPr as the last element."""
        sect = self.body.find(qn("w:sectPr"))
        if sect is not None:
            self.body.remove(sect)
            self.body.append(sect)

    # ----- headings
    def h1(self, text: str):
        self.chapter += 1
        self.fig = 0
        self.tbl = 0
        p = self.doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = True
        self._move_before_sect(p)
        return p

    def h1_unnumbered(self, text: str):
        p = self.doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = True
        return p

    def h2(self, text: str):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text: str):
        return self.doc.add_paragraph(text, style="Heading 3")

    def centered_title(self, text: str, page_break=True):
        p = self.doc.add_paragraph()
        if page_break:
            p.paragraph_format.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        p.paragraph_format.space_after = Pt(12)
        return p

    # ----- paragraphs
    def _inline(self, p, text: str, size=None, base_bold=False):
        # **bold**, *italic*, `code`, __underlined__ (primary keys)
        for i, part in enumerate(re.split(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*]+\*)", text)):
            if not part:
                continue
            if i % 2 == 0:
                # Plain text between the markers stays as it is, even when it starts with "*" (masked student IDs)
                r = p.add_run(part)
                if base_bold:
                    r.bold = True
            elif part.startswith("__"):
                r = p.add_run(part[2:-2])
                r.underline = True
                r.bold = True if base_bold else None
            elif part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                r.font.color.rgb = RGBColor.from_string(INLINE_CODE)
            else:
                r = p.add_run(part[1:-1])
                r.italic = True
                r.bold = base_bold or None
            if size:
                r.font.size = Pt(size)
        return p

    def p(self, text: str, indent=True, align="both", after=120):
        p = self.doc.add_paragraph()
        _para_format(p, after=after, first_line=284 if indent else 0, align=align)
        return self._inline(p, text)

    def note(self, text: str):
        """Note paragraph on a light background (notes, remarks)."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        _set_cell_shading(c, "EFF6FF")
        _set_cell_borders(c, "BFDBFE")
        _set_cell_margins(c, 80, 140, 80, 140)
        _set_cell_width(c, PAGE_WIDTH_TWIPS)
        p = c.paragraphs[0]
        _para_format(p, after=0, first_line=0, align="both")
        self._inline(p, text, size=12)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def bullets(self, items, level=0):
        for it in items:
            p = self.doc.add_paragraph(style="List Paragraph")
            _para_format(p, after=60, first_line=0, align="both")
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), str(level))
            numId = OxmlElement("w:numId")
            numId.set(qn("w:val"), "2")
            numPr.append(ilvl)
            numPr.append(numId)
            pPr.insert(1, numPr)
            self._inline(p, it)

    def numbered(self, items):
        """Manually numbered list 1), 2)... (Word would otherwise continue the numbering across lists)."""
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            _para_format(p, after=60, first_line=0, align="both")
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            self._inline(p, f"{i}) " + it)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True

    # ----- tables
    def table(self, headers, rows, widths_cm=None, caption: str | None = None, size=10.5, align=None,
              bold_first_col=False):
        if caption:
            self.tbl += 1
            cp = self.doc.add_paragraph(f"Bảng {self.chapter}.{self.tbl}. {caption}", style="TableCaption")
        ncol = len(headers)
        t = self.doc.add_table(rows=1, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths_cm is None:
            widths_cm = [16.0 / ncol] * ncol
        tw = [int(w * 567) for w in widths_cm]
        # header
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            _set_cell_shading(c, NAVY)
            _set_cell_borders(c)
            _set_cell_margins(c)
            _set_cell_width(c, tw[i])
            p = c.paragraphs[0]
            _para_format(p, after=0, line=252, first_line=0, align="center")
            r = p.add_run(str(h))
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            _cant_split(t.rows[-1])
            for i, val in enumerate(row):
                c = cells[i]
                _set_cell_borders(c)
                _set_cell_margins(c)
                _set_cell_width(c, tw[i])
                if ri % 2 == 1:
                    _set_cell_shading(c, "F5F8FC")
                lines = str(val if val is not None else "").split("\n")
                for li, line in enumerate(lines):
                    p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                    a = (align[i] if align else "left")
                    _para_format(p, after=0, line=252, first_line=0, align=a)
                    self._inline(p, line, size=size, base_bold=(bold_first_col and i == 0))
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        return t

    # ----- code boxes
    def code(self, title: str, code: str, lang: str = "sql", size=9.5):
        t = self.doc.add_table(rows=2, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.cell(0, 0), t.cell(1, 0)
        for c, fill in ((c0, CODE_TITLE_BG), (c1, CODE_BG)):
            _set_cell_shading(c, fill)
            _set_cell_borders(c)
            _set_cell_width(c, PAGE_WIDTH_TWIPS)
        _set_cell_margins(c0, 40, 140, 40, 140)
        _set_cell_margins(c1, 80, 140, 80, 140)
        _cant_split(t.rows[0])
        p = c0.paragraphs[0]
        _para_format(p, after=0, line=240, first_line=0, align="left", keep_next=True)
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        r.font.size = Pt(11)
        lines = code.expandtabs(4).split("\n")
        for li, line in enumerate(lines):
            p = c1.paragraphs[0] if li == 0 else c1.add_paragraph()
            _para_format(p, after=0, line=240, first_line=0, align="left")
            tokens = sql_tokens(line) if lang == "sql" else [(line, "000000")]
            if not line:
                tokens = [(" ", None)]
            for txt, color in tokens:
                r = p.add_run(txt)
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
                r._element.rPr.rFonts.set(qn("w:cs"), "Consolas")
                r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = RGBColor.from_string(color)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    # ----- figures
    def figure(self, path: Path, caption: str, width_cm: float = 15.5):
        _check_figure_caption(caption)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig += 1
        self.doc.add_paragraph(f"Hình {self.chapter}.{self.fig}. {caption}", style="FigureCaption")

    def figures_landscape(self, items, width_cm: float = 24.5):
        """Large figures (ERD, DFD...) on landscape pages (one figure per page), then back to portrait."""
        # Keep the portrait size BEFORE adding a section (the last section object changes after add_section)
        width, height = self.doc.sections[-1].page_width, self.doc.sections[-1].page_height
        landscape = self.doc.add_section(WD_SECTION.NEW_PAGE)
        landscape.orientation = WD_ORIENT.LANDSCAPE
        landscape.page_width, landscape.page_height = height, width
        for i, (path, caption) in enumerate(items):
            self.figure(path, caption, width_cm=width_cm)
            if i > 0:   # from the second figure on, the picture paragraph starts a new page
                picture = self.doc.paragraphs[-2]
                picture.paragraph_format.page_break_before = True
        portrait = self.doc.add_section(WD_SECTION.NEW_PAGE)
        portrait.orientation = WD_ORIENT.PORTRAIT
        portrait.page_width, portrait.page_height = width, height

    def figure_landscape(self, path: Path, caption: str, width_cm: float = 24.5):
        self.figures_landscape([(path, caption)], width_cm)

    # ----- tables of contents
    def toc(self, title: str, instr: str, placeholder: str, page_break=True):
        self.centered_title(title, page_break=page_break)
        p = self.doc.add_paragraph()
        _add_field(p, instr, placeholder)

    # ----- cover page: replace text inside the template
    def set_paragraph_text(self, index: int, text: str):
        par = self.doc.paragraphs[index]
        runs = par.runs
        if not runs:
            par.add_run(text)
            return
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""

    def fill_table(self, table_index: int, rows: list[list[str]], header: list[str] | None = None):
        """Overwrites a table of the template, adding/removing rows (copies the format of the first data row)."""
        t = self.doc.tables[table_index]
        if header:
            for i, h in enumerate(header):
                self._set_cell_text(t.rows[0].cells[i], h)
        data_rows = t.rows[1:]
        proto = copy.deepcopy(data_rows[0]._tr)
        while len(t.rows) - 1 < len(rows):
            t._tbl.append(copy.deepcopy(proto))
        while len(t.rows) - 1 > len(rows):
            t._tbl.remove(t.rows[-1]._tr)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                self._set_cell_text(t.rows[ri + 1].cells[ci], val)

    @staticmethod
    def _set_cell_text(cell, text: str):
        p = cell.paragraphs[0]
        runs = p.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
        else:
            p.add_run(text)
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)

    def enable_update_fields_on_open(self):
        settings = self.doc.settings.element
        el = settings.find(qn("w:updateFields"))
        if el is None:
            el = OxmlElement("w:updateFields")
            # CT_Settings order: updateFields comes before hdrShapeDefaults, footnotePr, compat, rsids...
            following = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                   "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                   "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                   "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
            anchor = next((c for c in settings if c.tag.split("}")[-1] in following), None)
            if anchor is not None:
                anchor.addprevious(el)
            else:
                settings.append(el)
        el.set(qn("w:val"), "true")

    def _normalize(self):
        """Reorders the children of w:tcPr and w:pPr as required by the OOXML schema."""
        order_by_element = {
            "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
                     "textDirection", "tcFitText", "vAlign", "hideMark"],
            "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
                    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
                    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
                    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
                    "sectPr", "pPrChange"],
        }
        for name, order in order_by_element.items():
            rank = {n: i for i, n in enumerate(order)}
            for el in self.body.iter(qn(f"w:{name}")):
                children = list(el)
                children.sort(key=lambda c: rank.get(c.tag.split("}")[-1], len(order)))
                for c in children:
                    el.remove(c)
                    el.append(c)

    def save(self, path: Path):
        self._normalize()
        self.doc.save(str(path))
